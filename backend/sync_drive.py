#!/usr/bin/env python3
"""Sync the public Innovaction Drive folder into the Spark knowledge corpus.

Usage:
  python sync_drive.py --folder-url 'https://drive.google.com/drive/folders/...'

The script downloads the shared folder, extracts readable text from PDF/TXT/MD/DOCX/PPTX,
creates provenance-preserving chunks, and writes data/chunks.json + data/corpus_manifest.json.
Images are kept in the manifest for the multimodal enrichment pass.
"""
from __future__ import annotations
import argparse, json, re, hashlib
from pathlib import Path
from datetime import datetime, timezone

ROOT=Path(__file__).resolve().parent.parent
CORPUS=ROOT/'data'/'drive_corpus'
OUT=ROOT/'data'/'chunks.json'
MANIFEST=ROOT/'data'/'corpus_manifest.json'

def clean(s:str)->str:
    return re.sub(r'\s+',' ',s or '').strip()

def chunk_text(text:str,size=1800,overlap=250):
    text=clean(text)
    if not text:return []
    out=[]; start=0
    while start<len(text):
        end=min(len(text),start+size); cut=text.rfind(' ',start,end)
        if cut<=start+size//2: cut=end
        out.append(text[start:cut]); start=max(cut-overlap,start+1)
    return out

def extract(path:Path):
    ext=path.suffix.lower(); pages=[]
    try:
        if ext=='.pdf':
            from pypdf import PdfReader
            for i,p in enumerate(PdfReader(str(path)).pages,1): pages.append((i,p.extract_text() or ''))
        elif ext in {'.txt','.md','.csv'}: pages=[(1,path.read_text(encoding='utf-8',errors='ignore'))]
        elif ext=='.docx':
            from docx import Document
            pages=[(1,'\n'.join(p.text for p in Document(str(path)).paragraphs))]
        elif ext=='.pptx':
            from pptx import Presentation
            for i,slide in enumerate(Presentation(str(path)).slides,1):
                pages.append((i,'\n'.join(sh.text for sh in slide.shapes if hasattr(sh,'text'))))
    except Exception as e:
        print(f'WARN extract {path.name}: {e}')
    return pages

def tags_for(text,name):
    hay=(name+' '+text).lower(); vocab={
      'innovation-games':['innovation game','juego de innovación','dinámica'],
      'design-thinking':['design thinking','empatía','ideation','ideación'],
      'canvas-value':['canvas','value proposition','propuesta de valor'],
      'customer-journey':['customer journey','journey map','cliente'],
      'trl':['trl','technology readiness','madurez tecnológica'],
      'poc-mvp':['poc','proof of concept','mvp','minimum viable'],
      'ip-patents':['patent','patente','propiedad intelectual','intellectual property'],
      'technology-watch':['vigilancia tecnológica','technology watch','competitive intelligence'],
      'metrics':['métrica','metric','kpi','indicador'],
      'sustainability':['sostenibilidad','sustainability','sustentabilidad'],
      'ai-industry40':['inteligencia artificial','artificial intelligence','industry 4.0','industria 4.0'],
      'strategy':['estrategia','strategy','roadmap']}
    return [tag for tag,keys in vocab.items() if any(k in hay for k in keys)]

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--folder-url',required=True); ap.add_argument('--skip-download',action='store_true'); a=ap.parse_args()
    CORPUS.mkdir(parents=True,exist_ok=True); OUT.parent.mkdir(parents=True,exist_ok=True)
    if not a.skip_download:
        import gdown
        print('Downloading/updating shared Drive corpus...')
        gdown.download_folder(url=a.folder_url,output=str(CORPUS),quiet=False,use_cookies=False,remaining_ok=True)
    chunks=[]; files=[]
    for path in sorted(p for p in CORPUS.rglob('*') if p.is_file()):
        rel=str(path.relative_to(CORPUS)); ext=path.suffix.lower(); sha=hashlib.sha256(path.read_bytes()).hexdigest()[:16]
        entry={'file':rel,'ext':ext,'bytes':path.stat().st_size,'sha256_16':sha,'kind':'image' if ext in {'.png','.jpg','.jpeg','.webp','.heic','.heif'} else 'document'}
        n=0
        for page,text in extract(path):
            for idx,part in enumerate(chunk_text(text)):
                chunks.append({'id':f'{sha}:{page}:{idx}','text':part,'source':rel,'page':page,'tags':tags_for(part,rel),'source_type':'drive'})
                n+=1
        entry['chunks']=n; files.append(entry)
    OUT.write_text(json.dumps(chunks,ensure_ascii=False,indent=2),encoding='utf-8')
    MANIFEST.write_text(json.dumps({'generated_at':datetime.now(timezone.utc).isoformat(),'folder_url':a.folder_url,'files':files,'chunks':len(chunks)},ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'Indexed {len(files)} files -> {len(chunks)} text chunks')
    print(f'Chunks: {OUT}\nManifest: {MANIFEST}')
if __name__=='__main__': main()
